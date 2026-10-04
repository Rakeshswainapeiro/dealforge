#!/usr/bin/env python3
"""Generate DealForge org data from the IHC subsidiary list.

Source: docs/List of Subsdiaries Q2 2026.xlsx  (sheet "Fulll list of subs")
Scope:  the Sirius International Holding Limited subtree ("keep Sirius for target now").
Outputs:
  - src/data/siriusOrgTree.js        (ES module: flat Organization rows + nested tree)
  - docs/DealForge_Org_Hierarchy_Sirius.md  (human-readable reference)
"""
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1] if (Path(__file__).resolve().parents[1] / "src").exists() else Path.cwd()
XLSX = ROOT / "docs" / "List of Subsdiaries Q2 2026.xlsx"
OUT_JS = ROOT / "src" / "data" / "siriusOrgTree.js"
OUT_MD = ROOT / "docs" / "DealForge_Org_Hierarchy_Sirius.md"

SHEET = "Fulll list of subs"
ROOT_NAME = "Sirius International Holding Limited"
REGION = range(944, 1189)  # rows bounding the Sirius subtree in the source sheet

HDR = re.compile(r"^Below (?:are the subsidiaries|is the subsidiary) of (.*?):?\s*$", re.I)


def norm(s: str) -> str:
    s = re.sub(r"\(.*?\)", "", str(s))
    s = s.replace("\u2013", "-").replace("\u2014", "-")
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def squash(s: str) -> str:
    """Aggressive match key: drop all punctuation/spacing (Quant Lase == Quantlase)."""
    return re.sub(r"[^a-z0-9]+", "", norm(s).lower())


_LEGAL = {"limited", "ltd", "llc", "rsc", "inc", "plc", "pjsc", "psc", "spc", "opc", "fz", "fze", "sa", "sarl", "ag", "bv", "llp"}


def corp(s: str) -> str:
    """Legal-suffix-insensitive key (Ltd == Limited)."""
    return " ".join(w for w in norm(s).split() if w not in _LEGAL)


def clean_name(s: str) -> str:
    s = re.sub(r"\s*\(note.*$", "", str(s), flags=re.I | re.S)  # drop "(note ...)" incl. nested parens
    s = re.sub(r"\s+", " ", s).strip().rstrip(",")
    return s


def slug(name: str) -> str:
    return "org-" + re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")[:60]


def infer_type(name: str, is_root: bool) -> str:
    n = name.lower()
    if is_root:
        return "platform"
    if "holding" in n or "holdings" in n:
        return "holding"
    if n.startswith("ihc "):
        return "holding"
    return "subsidiary"


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb[SHEET]

    # Source quirk: row 954 carries a stray "subsidiaries of Quantlase Lab LLC"
    # header inside Sirius's own child block; ignore it as a parent context.
    skip_headers = {norm("Quantlase Lab LLC")}

    nodes: dict[str, dict] = {}      # key = raw cleaned name (unique-ish)
    by_norm: dict[str, list[str]] = {}
    by_corp: dict[str, list[str]] = {}
    order: list[str] = []

    def ensure(raw: str) -> str | None:
        name = clean_name(raw)
        if not name:
            return None
        key = name
        if key not in nodes:
            nodes[key] = {
                "name": name,
                "place": None,
                "principal_activities": None,
                "ownership_pct": None,
                "note": None,
                "children": [],
            }
            by_norm.setdefault(norm(name), []).append(key)
            by_corp.setdefault(corp(name), []).append(key)
            order.append(key)
        return key

    # 1. register every entity row in the region
    for r in REGION:
        raw = ws.cell(r, 2).value
        if not raw:
            continue
        s = str(raw).strip()
        if s.startswith("Below") or s in ("Name of subsidiary", "S.No"):
            continue
        key = ensure(raw)
        if not key:
            continue
        node = nodes[key]
        if ws.cell(r, 3).value:
            node["place"] = str(ws.cell(r, 3).value).strip()
        if ws.cell(r, 4).value:
            node["principal_activities"] = str(ws.cell(r, 4).value).strip()
        if ws.cell(r, 5).value is not None:
            node["ownership_pct"] = ws.cell(r, 5).value
        if ws.cell(r, 8).value:
            node["note"] = str(ws.cell(r, 8).value).strip()

    root_key = ensure(ROOT_NAME)

    def find_header_parent(label: str) -> str | None:
        # Exact (normalised) name match first; fall back to legal-suffix-insensitive
        # matching (Ltd == Limited). Sirius's direct children are listed in one block
        # interrupted by a stray header, so fuzzy token matching is avoided.
        for table in (by_norm, by_corp):
            key = norm(label) if table is by_norm else corp(label)
            cand = table.get(key, [])
            for k in cand:
                if k != root_key:
                    return k
            if cand:
                return cand[0]
        return None

    def is_ancestor(descendant_key: str, candidate_key: str, parent: dict[str, str]) -> bool:
        cur = parent.get(descendant_key)
        seen = set()
        while cur and cur not in seen:
            if cur == candidate_key:
                return True
            seen.add(cur)
            cur = parent.get(cur)
        return False

    parent_of: dict[str, str] = {}
    cur: str | None = None

    for r in REGION:
        raw = ws.cell(r, 2).value
        if not raw:
            continue
        s = str(raw).strip()
        m = HDR.match(s)
        if m:
            label = m.group(1).strip()
            if norm(label) in skip_headers:
                continue
            cur = find_header_parent(label)
            if cur is None:
                # header-only parent (never listed as a child row) -> create + nest under root
                cur = ensure(label)
                if cur and cur != root_key and cur not in parent_of:
                    parent_of[cur] = root_key
                    nodes[root_key]["children"].append(cur)
            continue
        if s.startswith("Below") or s in ("Name of subsidiary", "S.No"):
            continue
        key = ensure(raw)
        if not key or key == root_key or key == cur:
            continue
        parent = cur or root_key
        if key in parent_of:
            continue
        if is_ancestor(parent, key, parent_of):  # cycle guard
            continue
        parent_of[key] = parent
        nodes[parent]["children"].append(key)

    # 2. attach any nodes never placed (orphan header-parents) to the root
    for key in order:
        if key != root_key and key not in parent_of:
            parent_of[key] = root_key
            nodes[root_key]["children"].append(key)

    # 3. flatten with ids / type / level
    ids: dict[str, str] = {}
    used: dict[str, int] = {}
    for key in order:
        base = slug(nodes[key]["name"])
        n = used.get(base, 0)
        used[base] = n + 1
        ids[key] = base if n == 0 else f"{base}-{n + 1}"

    def level_of(key: str) -> int:
        lvl, cur = 0, parent_of.get(key)
        seen = set()
        while cur and cur not in seen:
            lvl += 1
            seen.add(cur)
            cur = parent_of.get(cur)
        return lvl

    flat = []
    for key in order:
        node = nodes[key]
        flat.append({
            "id": ids[key],
            "name": node["name"],
            "type": infer_type(node["name"], key == root_key),
            "parent_id": ids[parent_of[key]] if key in parent_of else None,
            "level": level_of(key),
            "ownership_pct": node["ownership_pct"],
            "geography": node["place"],
            "principal_activities": node["principal_activities"],
            "note": node["note"],
            "source": "List of Subsdiaries Q2 2026.xlsx",
            "visibility": "organization",
        })

    # nested tree (for convenient rendering)
    def build(key: str) -> dict:
        n = next(x for x in flat if x["id"] == ids[key])
        return {
            "id": n["id"],
            "name": n["name"],
            "type": n["type"],
            "ownership_pct": n["ownership_pct"],
            "geography": n["geography"],
            "children": [build(c) for c in nodes[key]["children"]],
        }

    tree = build(root_key)

    # 4. write the ES module
    OUT_JS.parent.mkdir(parents=True, exist_ok=True)
    OUT_JS.write_text(
        "// AUTO-GENERATED from docs/List of Subsdiaries Q2 2026.xlsx (sheet \"Fulll list of subs\").\n"
        "// Scope: the Sirius International Holding Limited subtree (\"keep Sirius for target now\").\n"
        "// Regenerate with: python3 scripts/generate_org_tree.py\n"
        "/* eslint-disable */\n\n"
        f"export const SIRIUS_ORG_ROOT_ID = {json.dumps(ids[root_key])};\n\n"
        f"export const SIRIUS_ORGANIZATIONS = {json.dumps(flat, indent=2, ensure_ascii=False)};\n\n"
        f"export const SIRIUS_ORG_TREE = {json.dumps(tree, indent=2, ensure_ascii=False)};\n\n"
        "export default SIRIUS_ORGANIZATIONS;\n"
    )

    # 5. write the markdown reference
    lines = [
        "# DealForge — Sirius Organisation Hierarchy (reference)",
        "",
        "**Status:** Extracted reference (v1.0) from `docs/List of Subsdiaries Q2 2026.xlsx`",
        "**Scope:** the **Sirius International Holding Limited** subtree — Sirius is the target root for now.",
        "**Source sheet:** `Fulll list of subs` · **Rows:** Sirius block ~944–1188 · **Ownership:** column E (2026).",
        "",
        "> The source is a flat DFS listing; parentage is implied by `Below are the subsidiaries of <X>`",
        "> headers. Names normalise to unique ids in `src/data/siriusOrgTree.js`.",
        "",
        f"**Entities in scope:** {len(flat)}  ·  **Root:** {ROOT_NAME} (ownership in IHC: "
        f"{next(x for x in flat if x['id']==ids[root_key])['ownership_pct']})",
        "",
        "## Direct holdings of Sirius",
        "",
    ]
    root_children = nodes[root_key]["children"]
    for c in root_children:
        kids = len(nodes[c]["children"])
        lines.append(f"- **{nodes[c]['name']}** — {kids} direct children" if kids else f"- **{nodes[c]['name']}**")
    lines += ["", "## Full tree", "", "```"]

    def md(key: str, depth: int) -> None:
        node = nodes[key]
        own = f"  [{node['ownership_pct']}]" if node["ownership_pct"] not in (None, "") else ""
        lines.append("  " * depth + "- " + node["name"] + own)
        for c in node["children"]:
            md(c, depth + 1)

    md(root_key, 0)
    lines += ["```", "",
              "## Notes & inferred parentage", "",
              "- Parent links come from the file's section headers. One parent — **International Tech Group – Sole Proprietorship LLC** — appears only as a header (never as a child row), so it is created and nested under Sirius; its four children (Emircom, Data Center Vault, Sand Dynamics, Cyber Gate Defense) follow from there.",
              "- A stray *\"subsidiaries of Quantlase Lab LLC\"* header (row 954) sits inside Sirius's own child block and is ignored, so Sanimed, CMC Holding, Esyasoft Holding and QData remain direct children of Sirius.",
              "- Owners whose names differ only by legal suffix are matched (`Esyasoft Enterprise Holding RSC Ltd` ≈ `…RSC Limited`).",
              "- `Cyber Gate Defense LLC (Dubai)` is kept distinct from its parent `Cyber Gate Defense LLC`.",
              "- Ownership % is taken from the 2026 column; `null` where the file leaves it blank.",
              "- This file is generated; edit the generator, not the output.",
              ""]
    OUT_MD.write_text("\n".join(lines))

    print(f"nodes={len(flat)}  root={ids[root_key]}")
    print("direct children:", ", ".join(nodes[c]["name"] for c in root_children))
    print("wrote", OUT_JS)
    print("wrote", OUT_MD)


if __name__ == "__main__":
    main()
